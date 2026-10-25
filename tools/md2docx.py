"""Minimal markdown -> docx (headings, paragraphs, bullets/numbered, tables kept together, images, **bold**/*italic*/links)."""
import re, sys, os
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
def runs(p, text):
    text=re.sub(r'\[([^\]]+)\]\((http[^)]+)\)', r'\1 (\2)', text)
    for part in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not part: continue
        if part.startswith('**'): p.add_run(part[2:-2]).bold=True
        elif part.startswith('*') and len(part)>1: p.add_run(part[1:-1]).italic=True
        else: p.add_run(part)
def keep(p):
    pPr=p._p.get_or_add_pPr()
    for t in ('w:keepNext','w:keepLines'): pPr.append(OxmlElement(t))
def convert(md, out, landscape=False, base=None):
    base=base or os.path.dirname(os.path.abspath(md))
    d=Document(); st=d.styles['Normal']; st.font.name='Arial'; st.font.size=Pt(9.5)
    s=d.sections[0]
    if landscape:
        s.orientation=WD_ORIENT.LANDSCAPE; s.page_width, s.page_height = s.page_height, s.page_width
    for a in ('left_margin','right_margin','top_margin','bottom_margin'): setattr(s,a,Cm(1.5))
    width=s.page_width-s.left_margin-s.right_margin
    L=open(md).read().split('\n'); i=0
    while i<len(L):
        l=L[i]
        if not l.strip(): i+=1; continue
        m=re.match(r'^(#{1,4})\s+(.*)',l)
        if m: keep(d.add_heading(re.sub(r'\*','',m.group(2)), level=len(m.group(1)))); i+=1; continue
        m=re.match(r'^!\[[^\]]*\]\(([^)]+)\)',l.strip())
        if m:
            p=d.add_picture(os.path.join(base,m.group(1)), width=int(width*0.85)); i+=1; continue
        if l.lstrip().startswith('|'):
            rows=[]
            while i<len(L) and L[i].lstrip().startswith('|'):
                c=[x.strip() for x in L[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',x) for x in c if x): rows.append(c)
                i+=1
            n=max(map(len,rows)); t=d.add_table(rows=len(rows), cols=n); t.style='Table Grid'
            fs=Pt(7 if n>8 else 8)
            for r,row in enumerate(rows):
                tr=t.rows[r]._tr; trPr=tr.get_or_add_trPr(); cs=OxmlElement('w:cantSplit'); trPr.append(cs)
                if r==0: h=OxmlElement('w:tblHeader'); trPr.append(h)
                for c,v in enumerate(row):
                    p=t.cell(r,c).paragraphs[0]; runs(p, v if r else f'**{v.strip("*")}**' if v else '')
                    for rn in p.runs: rn.font.size=fs
                    if r<len(rows)-1: keep(p)
            d.add_paragraph(); continue
        m=re.match(r'^(\s*)([-*]|\d+\.)\s+(.*)',l)
        if m:
            sty='List Number' if m.group(2)[0].isdigit() else 'List Bullet'
            if len(m.group(1))>=2: sty+=' 2'
            p=d.add_paragraph(style=sty); runs(p,m.group(3)); i+=1; continue
        buf=[l.strip()]; i+=1
        while i<len(L) and L[i].strip() and not re.match(r'^(#|\||!\[|\s*([-*]|\d+\.)\s)',L[i]): buf.append(L[i].strip()); i+=1
        p=d.add_paragraph(); runs(p,' '.join(buf))
        j=i
        while j<len(L) and not L[j].strip(): j+=1
        if j<len(L) and len(p.text)<200 and (re.match(r'^\s*(\||!\[)',L[j]) or (L[j].startswith('*') and len(L[j])<120)): keep(p)
    d.save(out)
if __name__=='__main__': convert(sys.argv[1], sys.argv[2], landscape='--landscape' in sys.argv)
