"""Simple class-level US PD eligible-pool model (published inputs). Writes pd_simple2_model.xlsx/json."""
import json, math
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
L=dict(marras="https://doi.org/10.1038/s41531-018-0058-0",macleod="https://doi.org/10.1002/mds.25898",
 aarsland="https://doi.org/10.1002/mds.20527",gibson="https://doi.org/10.1002/mds.29918",
 track="https://pmc.ncbi.nlm.nih.gov/articles/PMC11145112/",camp="https://pmc.ncbi.nlm.nih.gov/articles/PMC7063985/",
 rizzo="https://doi.org/10.1212/WNL.0000000000002350",venk="https://pmc.ncbi.nlm.nih.gov/articles/PMC10898494/")
# (key, label, value or formula, tag, source)
I=[("P20","US PD prevalence 2020 (age 45+, minimum)",930000,"A","Marras 2018 "+L['marras']),
 ("P30","US PD prevalence 2030 (age 45+, minimum)",1238000,"A","Marras 2018 "+L['marras']),
 ("G","Growth of prevalence and entrants, per year","=(B3/B2)^(1/10)-1","C","log-linear between Marras 2020 and 2030"),
 ("P26","US PD prevalence FY2026","=B2*(1+B4)^6","C","Marras 2020 grown 6 years at g"),
 ("D","Death rate in PD, per year",0.05,"A","Macleod 2014: survival falls ~5% per year of follow-up "+L['macleod']),
 ("DEM","Dementia point prevalence in PD",0.28,"A/S","Aarsland 2005: 24.5% (all studies) to 31.1% (best studies); midpoint "+L['aarsland']),
 ("S4","Share of prevalent PD diagnosed >=4 years","=EXP(-4*(B6+B4))","C","survival e^(-d t) weighted for growth g; check: prevalence/incidence 1.1M/90k = 12 yr mean"),
 ("MC","Motor complications among non-demented >=4 yrs",0.5,"S","anchored: off periods 35% at 4-6 yrs, 59% at 8-10 yrs "+L['track']+"; fluctuations 54% at 5 yrs "+L['camp']),
 ("DX","Confirmed levodopa-responsive idiopathic PD",0.806,"A","Rizzo 2016 clinical-diagnosis accuracy (proxy) "+L['rizzo']),
 ("FIT","Medically fit for neurosurgery",0.9,"S","judgment; no published population figure"),
 ("INC","Dementia incidence in PD, per year",0.045,"A","Gibson 2024 meta-analysis: 4.45 per 100 person-years "+L['gibson']),
 ("M","Leave rate m (death + dementia)","=B6+B12","C","sum of the two rates"),
 ("DBS","DBS procedures for PD per year (held flat)",3990,"A","Venkatraman 2024, NIS 2019 "+L['venk']),
 ("U","Share of pool treated per year from launch",0.07,"S","David May's proposal"),
 ("LAU","Launch fiscal year",2029,"S","assumption"),
 ("PR","Price per patient, $ (illustration)",20000,"S","assumption")]
R={k:i+2 for i,(k,*_) in enumerate(I)}
def v(k):
    x={k2:val for k2,_,val,*_ in I}
    return x[k]
X={}
X['P20'],X['P30']=930000,1238000; X['G']=(X['P30']/X['P20'])**0.1-1; X['P26']=X['P20']*(1+X['G'])**6
for k in ('D','DEM','MC','DX','FIT','INC','DBS','U','LAU','PR'): X[k]=v(k)
X['S4']=math.exp(-4*(X['D']+X['G'])); X['M']=X['D']+X['INC']
c=dict(nd4=X['S4']-X['DEM']); c['a']=X['P26']*c['nd4']; c['b']=c['a']*X['MC']; c['c']=c['b']*X['DX']; c['elig']=c['c']*X['FIT']
c['prior']=X['DBS']/X['M']; c['U0']=c['elig']-c['prior']; c['E0']=(X['M']+X['G'])*c['U0']+X['DBS']
c['Pss']=X['U']/(X['M']+X['U']+X['G']); c['Plife']=X['U']/(X['M']+X['U'])
rows=[]; S=c['U0']
for t in range(2026,2046):
    E=c['E0']*(1+X['G'])**(t-2026); Lv=X['M']*S; T=X['DBS'] if t<X['LAU'] else X['U']*S
    rows.append(dict(t=t,S=S,E=E,L=Lv,T=T,end=S+E-Lv-T,ratio=T/E,rev=T*X['PR']/1e6)); S=S+E-Lv-T
json.dump(dict(inputs=X,chain=c,rows=rows),open('pd_simple2_model.json','w'),indent=1)
# workbook
wb=Workbook(); B=Font(bold=True); H=PatternFill('solid',fgColor='DDEBF7')
ws=wb.active; ws.title='Method'
for i,s in enumerate(["US class-level eligible pool for a one-time PD therapy; single base case; FY2026-45.",
 "Start: published US PD prevalence (Marras 2018, minimum estimates) grown log-linearly to FY2026.",
 "Screens: diagnosed >=4 yrs and no dementia (joint), motor complications, levodopa-responsive, fit. No age cap.",
 "Previously DBS-treated patients removed: DBS per year / m (rough steady-state approximation).",
 "Each year: Pool end = Pool start + Entrants - Leavers - Treated.",
 "Leavers = m x pool start; m = death rate + dementia incidence (published).",
 "Entrants = (m + g) x starting pool + DBS, so the pool would grow at g with no new therapy; then grow at g.",
 "Treated = DBS only before launch; from FY2029 a fixed 7% of the pool start.",
 "Steady state: treated / entrants -> u/(m+u+g); a cohort's lifetime share is u/(m+u).",
 "Tags: A published, C calculated, S judgment. Every input with its source is on the Inputs tab."],1):
    ws.cell(i,1,s)
ws.column_dimensions['A'].width=110
wi=wb.create_sheet('Inputs'); 
for j,h in enumerate(['Input','Value','Tag','Source'],1): x=wi.cell(1,j,h); x.font=B; x.fill=H
for k,lab,val,tag,src in I:
    r=R[k]; wi.cell(r,1,lab); wi.cell(r,2,val); wi.cell(r,3,tag); wi.cell(r,4,src)
    fmt={'G':'0.0%','D':'0.0%','DEM':'0%','S4':'0%','MC':'0%','DX':'0.0%','FIT':'0%','INC':'0.0%','M':'0.000','U':'0%','LAU':'0'}.get(k,'#,##0')
    wi.cell(r,2).number_format=fmt
for col,w in zip('ABCD',(46,14,6,120)): wi.column_dimensions[col].width=w
wf=wb.create_sheet('Flow'); r=lambda k:f"Inputs!$B${R[k]}"
chain=[("US PD prevalence FY2026",f"={r('P26')}"),
 ("x (diagnosed >=4 yrs minus dementia) = no dementia and >=4 yrs",f"=B2*({r('S4')}-{r('DEM')})"),
 ("x motor complications",f"=B3*{r('MC')}"),("x levodopa-responsive",f"=B4*{r('DX')}"),
 ("x medically fit = medically eligible",f"=B5*{r('FIT')}"),("- previously DBS-treated (DBS / m)",f"=-{r('DBS')}/{r('M')}"),
 ("Untreated eligible pool, start FY2026","=B6+B7"),("Entrants FY2026 = (m+g) x pool + DBS",f"=({r('M')}+{r('G')})*B8+{r('DBS')}"),
 ("Eligible share of PD","=B6/B2"),("Steady-state share of entrants treated u/(m+u+g)",f"={r('U')}/({r('M')}+{r('U')}+{r('G')})"),
 ("Cohort lifetime share treated u/(m+u)",f"={r('U')}/({r('M')}+{r('U')})")]
wf.cell(1,1,'Starting pool').font=B
for i,(lab,f) in enumerate(chain,2):
    wf.cell(i,1,lab); wf.cell(i,2,f); wf.cell(i,2).number_format='0.0%' if i>=11 else '#,##0'
h0=15
for j,hd in enumerate(['FY','Pool start','Entrants','Leavers','Treated','Pool end','Treated/Entrants','Revenue $M'],1):
    x=wf.cell(h0,j,hd); x.font=B; x.fill=H
for k,t in enumerate(range(2026,2046)):
    q=h0+1+k; wf.cell(q,1,t)
    wf.cell(q,2,"=B8" if k==0 else f"=F{q-1}")
    wf.cell(q,3,f"=$B$9*(1+{r('G')})^(A{q}-2026)")
    wf.cell(q,4,f"={r('M')}*B{q}")
    wf.cell(q,5,f"=IF(A{q}<{r('LAU')},{r('DBS')},{r('U')}*B{q})")
    wf.cell(q,6,f"=B{q}+C{q}-D{q}-E{q}"); wf.cell(q,7,f"=E{q}/C{q}"); wf.cell(q,8,f"=E{q}*{r('PR')}/1e6")
    for j in range(2,7): wf.cell(q,j).number_format='#,##0'
    wf.cell(q,7).number_format='0.00'; wf.cell(q,8).number_format='#,##0'
wf.column_dimensions['A'].width=58
for col in 'BCDEFGH': wf.column_dimensions[col].width=14
wb.save('pd_simple2_model.xlsx')
print({k:round(val,4) if isinstance(val,float) else val for k,val in X.items()}); print({k:round(val,3) for k,val in c.items()})
for x in rows: print(x['t'],round(x['S']),round(x['E']),round(x['T']),round(x['ratio'],3))
