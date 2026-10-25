"""Simple single-scenario eligible-pool model (class level). Cohort numbers come from pd_pool_model.cohort() with label-style screens
(any motor complication, levodopa-responsive, fit, age-specific dementia, >=4 yrs, no age cap). Writes pd_simple_model.json and pd_simple_model.xlsx."""
import json, pd_pool_model as M
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
p=dict(M.BASE,dis=1.0,cap=100); C=M.cohort(p); R={r['t']:r for r in C['rows']}; b=C['snaps']['2026']['bands']
PD=b and C['snaps']['2026']['PD']; NODEM=sum(x['nodem'] for x in b.values()); MC=sum(x['mc'] for x in b.values())
m=round((R[2026]['death']+R[2026]['dem'])/R[2025]['E'],2)      # exits / pool start
g=round(M.G,3)
X=dict(PD=round(PD,-3),s_alive=round(NODEM/PD,3),s_mc=round(MC/NODEM,3),dx=0.806,fit=0.90,N26=round(R[2026]['N'],-2),g=g,m=m,
       r=0.55,w=0.60,c=0.90,dbs=3990,launch=2029,price=20000,E19=R[2019]['E'],N19=R[2019]['N'])
def sim(X):
    P=X['r']*X['w']*X['c']; u=P*(X['m']+X['g'])/(1-P)
    gross=X['PD']*X['s_alive']*X['s_mc']*X['dx']*X['fit']; prior=X['dbs']/X['m']; S=gross-prior; rows=[]
    for t in range(2026,2046):
        E=X['N26']*(1+X['g'])**(t-2026); L=X['m']*S; T=X['dbs'] if t<X['launch'] else u*S
        rows.append(dict(t=t,S=S,E=E,L=L,T=T,end=S+E-L-T,ratio=T/E)); S=S+E-L-T
    return dict(P=P,u=u,gross=gross,prior=prior,rows=rows)
if __name__=='__main__':
    o=sim(X); json.dump(dict(inputs=X,**o,cohort_check=dict(E2019=R[2019]['E'],N2019=R[2019]['N'],E2026=R[2026]['E'],N2026=R[2026]['N'],
        N2045=R[2045]['N'],mu=R[2026]['mu'],exits26=R[2026]['death']+R[2026]['dem'])),open('pd_simple_model.json','w'),indent=1)
    print('P',o['P'],'u',round(o['u'],4),'gross',round(o['gross']),'prior',round(o['prior']),'start',round(o['rows'][0]['S']))
    for r in o['rows']: print(r['t'],*(round(r[k]) for k in ('S','E','L','T','end')),round(r['ratio'],3))
    # long-run convergence check (hold everything, run to 2100)
    S=o['rows'][-1]['end']; E=o['rows'][-1]['E']
    for _ in range(55): E*=1+X['g']; T=o['u']*S; ratio=T/E; S=S+E-X['m']*S-T
    print('2100 T/E',round(ratio,4),'P',o['P'])
    # ---------- xlsx
    wb=Workbook(); B=Font(bold=True); I=PatternFill('solid',fgColor='FFF2CC'); H=PatternFill('solid',fgColor='DDEBF7'); wrap=Alignment(wrap_text=True,vertical='top')
    ws=wb.active; ws.title='Method'
    for i,s in enumerate(["US eligible pool for a one-time restorative PD therapy: class level, one scenario (base case for agreement)",
     "1. Medically eligible = label-style criteria: levodopa-responsive PD, ≥4 years, any motor complication, no dementia, medically fit, no age cap.",
     "2. Funnel: medically eligible > referred > willing > covered > treated. P = referral × willing × covered (judgment, anchored in pd_research_notes.md).",
     "3. Pool at FY2026 start = eligible patients from the cohort model (screens on the Flow tab) less patients already treated with DBS.",
     "4. Each year: Pool end = Pool start + Entrants − Leavers − Treated.",
     "5. Entrants = newly eligible patients (cohort model), growing with the ageing population. Leavers = m × Pool start (death or dementia).",
     "6. Before launch (FY2029) only DBS is treated (3,990 a year, held flat). From launch, Treated = u × Pool start.",
     "7. u is set so that, once the backlog is used up, Treated = P × Entrants: u = P × (m + g) ÷ (1 − P), where g is entrant growth.",
     "8. Treated = class demand. It is not split between therapies, and it is not limited by capacity or competition.",
     "Yellow = input. Tags: A published, C calculated, S judgment (★), P instruction. Sources on the Inputs tab."],1):
        ws.cell(i,1,s).font=B if i==1 else Font()
    ws.column_dimensions['A'].width=125
    wi=wb.create_sheet('Inputs')
    for j,h in enumerate(['Input','Value','Tag','Source / rationale'],1): c=wi.cell(1,j,h); c.font=B; c.fill=H
    INP=[('PD prevalent, US, FY2026',X['PD'],'#,##0','C','Cohort model: 90k diagnoses/yr (Parkinson\'s Foundation), Hirsch incidence, SSA life table × 1.52 mortality ratio, Census 2024'),
     ('Share alive without dementia and ≥4 yrs since diagnosis',X['s_alive'],'0.00','C','Cohort model; dementia CamPaIGN 46% at 10 yrs (≥65 at diagnosis), Sydney 48% at 15 yrs (<65)'),
     ('… of whom with motor complications',X['s_mc'],'0.00','C','Cohort model; Ahlskog 2001 (~40% by 4–6 yrs), Hely 2005 (~95% by 15 yrs)'),
     ('Levodopa-responsive idiopathic PD',X['dx'],'0.00','A','Rizzo 2016, Neurology: diagnostic accuracy 80.6%'),
     ('Medically fit for surgery',X['fit'],'0.00','S ★','Judgment, range 70–90%'),
     ('Newly eligible patients (entrants), FY2026',X['N26'],'#,##0','C','Cohort model, same screens (21.6k in 2019)'),
     ('Entrant growth per year (g)',X['g'],'0.0%','C','Census 2020–24, age-weighted PD incidence growth'),
     ('Annual leave rate from the pool (m)',X['m'],'0%','C','Cohort model: (deaths + dementia onsets) ÷ pool start, FY2026 = 21.5k ÷ 162.8k'),
     ('Referral × Willing × Covered',f"{X['r']} × {X['w']} × {X['c']}",'@','S ★','Referral 0.55 (30–80%: PARADISE, Olmsted); willing 0.60 (35–43% consent to referral, Dinkelbach 2017; 58.7% implanted after evaluation); covered 0.90 (Medicare covers label)'),
     ('DBS for PD per year (2019, held flat)',X['dbs'],'#,##0','A','Venkatraman 2024, NIS 2019, inpatient, all ages'),
     ('Launch of first one-time therapy',X['launch'],'0','S ★','Judgment'),
     ('Revenue per patient (illustration)',X['price'],'$#,##0','P','David May')]
    REF={}
    for i,(n,v,f,tag,src) in enumerate(INP,2):
        wi.cell(i,1,n); c=wi.cell(i,2,v); c.fill=I; c.number_format=f; wi.cell(i,3,tag); wi.cell(i,4,src).alignment=wrap; REF[i]=f'Inputs!$B${i}'
    # r,w,c separate hidden cells? keep explicit: put factors in E..G of row 10
    wi.cell(1,5,'Referral').font=B; wi.cell(1,6,'Willing').font=B; wi.cell(1,7,'Covered').font=B
    for j,v in zip((5,6,7),(X['r'],X['w'],X['c'])): c=wi.cell(10,j,v); c.fill=I; c.number_format='0.00'
    wi.cell(10,2,'=E10*F10*G10'); wi.cell(10,2).number_format='0.00'; wi.cell(10,2).fill=PatternFill()
    for col,w in zip('ABCDEFG',[48,12,7,95,10,10,10]): wi.column_dimensions[col].width=w
    PDc,SA,SM,DX,FIT,N26,G,Mc,Pc,DBS,LAU,PR=[REF[i] for i in range(2,14)]
    wf=wb.create_sheet('Flow')
    wf.cell(1,1,'Starting pool, FY2026 (label-style screens)').font=B
    chain=[('PD prevalent',f'={PDc}'),('× alive, no dementia, ≥4 yrs',f'=B2*{SA}'),('× with motor complications',f'=B3*{SM}'),('× levodopa-responsive',f'=B4*{DX}'),
           ('× medically fit = medically eligible',f'=B5*{FIT}'),('− previously DBS-treated still in pool (DBS ÷ m) ★',f'=-{DBS}/{Mc}'),('= Untreated eligible pool, start FY2026',f'=B6+B7')]
    for i,(n,f) in enumerate(chain,2):
        wf.cell(i,1,n); c=wf.cell(i,2,f); c.number_format='#,##0,"k"'
    wf.cell(8,2).font=B; wf.cell(8,1).font=B
    wf.cell(10,1,'Derived').font=B
    wf.cell(11,1,'P = referral × willing × covered'); wf.cell(11,2,f'={Pc}'); wf.cell(11,2).number_format='0.00'
    wf.cell(12,1,'u = P × (m + g) ÷ (1 − P), yearly share of the pool treated'); wf.cell(12,2,f'=B11*({Mc}+{G})/(1-B11)'); wf.cell(12,2).number_format='0.0%'
    wf.cell(13,1,'Sanity check: DBS-implied P = DBS ÷ entrants (2019: 3,990 ÷ 21.6k)'); wf.cell(13,2,f'={DBS}/{round(X["N19"],-2)}'); wf.cell(13,2).number_format='0.00'
    wf.cell(15,1,'Yearly flow (patients)').font=B
    hd=['FY','Pool start','Entrants','Leavers','Treated','Pool end','Treated ÷ Entrants','Revenue at $20k ($M)']
    for j,h in enumerate(hd,1): c=wf.cell(16,j,h); c.font=B; c.fill=H; c.alignment=wrap
    for k,t in enumerate(range(2026,2046)):
        r=17+k; wf.cell(r,1,t)
        wf.cell(r,2,'=B8' if k==0 else f'=F{r-1}'); wf.cell(r,3,f'={N26}*(1+{G})^(A{r}-2026)'); wf.cell(r,4,f'={Mc}*B{r}')
        wf.cell(r,5,f'=IF(A{r}<{LAU},{DBS},$B$12*B{r})'); wf.cell(r,6,f'=B{r}+C{r}-D{r}-E{r}'); wf.cell(r,7,f'=E{r}/C{r}'); wf.cell(r,8,f'=E{r}*{PR}/1e6')
        for j in range(2,7): wf.cell(r,j).number_format='#,##0.0,"k"'
        wf.cell(r,7).number_format='0.00'; wf.cell(r,8).number_format='0'
    wf.column_dimensions['A'].width=62
    for col in 'BCDEFGH': wf.column_dimensions[col].width=13
    wf.freeze_panes='A17'
    wb.save('pd_simple_model.xlsx'); print('saved')
